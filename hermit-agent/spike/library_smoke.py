"""Phase 3.5 study, part 2: does each candidate work in the app's Pyodide?

Downloads every candidate from spike/out/library_study.json (plus the dependencies
Pyodide lacks) to spike/out/wheels/, opens the built app in headless Chromium, uploads
the wheels to the workspace, unpacks each one into site-packages (what bundling would
do; a source-only package is unpacked from its sdist), and runs one small real task per
library through the app's own runInWorker. Pyodide's own packages that a candidate
needs (lxml, Pillow, …) are loaded by importing them first, as the harness would.

    python3 build.py && ../benchmark/.venv/bin/python spike/library_smoke.py

Results: spike/out/library_smoke.json and a table.
"""
import hashlib
import io
import json
import pathlib
import re
import sys
import tarfile
import time
import urllib.request

from playwright.sync_api import sync_playwright

sys.stdout.reconfigure(line_buffering=True)
ROOT = pathlib.Path(__file__).resolve().parent.parent
STUDY = ROOT / "spike" / "out" / "library_study.json"
WHEELS = ROOT / "spike" / "out" / "wheels"
OUT = ROOT / "spike" / "out" / "library_smoke.json"
APP = ROOT / "dist" / "hermit-agent-standalone.html"
LOCK = ROOT / "libs" / "pyodide-0.29.5" / "pyodide-lock.json"

# One small real task per library. Each prints "OK <detail>" when it works.
SMOKE = {
    "openpyxl": '''
import openpyxl, pandas as pd
wb = openpyxl.Workbook(); ws = wb.active; ws.append(["month", "total"]); ws.append(["2026-01", 200.5]); ws["C2"] = "=B2*2"; wb.save("t.xlsx")
df = pd.read_excel("t.xlsx"); pd.DataFrame({"a": [1, 2]}).to_excel("p.xlsx", index=False)
print("OK", list(df.columns), pd.read_excel("p.xlsx")["a"].tolist())''',
    "xlsxwriter": '''
import xlsxwriter, pandas as pd
wb = xlsxwriter.Workbook("c.xlsx"); ws = wb.add_worksheet(); ws.write_column("A1", [1, 2, 3])
ch = wb.add_chart({"type": "line"}); ch.add_series({"values": "=Sheet1!$A$1:$A$3"}); ws.insert_chart("C1", ch); wb.close()
pd.DataFrame({"a": [1]}).to_excel("x.xlsx", engine="xlsxwriter")
print("OK", __import__("os").path.getsize("c.xlsx"))''',
    "python-docx": '''
import docx
d = docx.Document(); d.add_heading("Report", 1); d.add_paragraph("Umlauts äöü")
t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Item"; t.cell(1, 0).text = "nuts"; d.save("r.docx")
print("OK", [p.text for p in docx.Document("r.docx").paragraphs])''',
    "docx2txt": '''
import docx, docx2txt
d = docx.Document(); d.add_paragraph("hello docx2txt"); d.save("a.docx")
print("OK", docx2txt.process("a.docx").strip())''',
    "mammoth": '''
import docx, mammoth
d = docx.Document(); d.add_heading("Title", 1); d.add_paragraph("body"); d.save("m.docx")
print("OK", mammoth.convert_to_html(open("m.docx", "rb")).value[:60])''',
    "docxtpl": '''
import docx
from docxtpl import DocxTemplate
d = docx.Document(); d.add_paragraph("Hello {{ name }}"); d.save("tpl.docx")
t = DocxTemplate("tpl.docx"); t.render({"name": "Ada"}); t.save("out.docx")
print("OK", docx.Document("out.docx").paragraphs[0].text)''',
    "python-pptx": '''
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Q1"
cd = CategoryChartData(); cd.categories = ["a", "b"]; cd.add_series("s", (1, 2))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd); p.save("d.pptx")
print("OK", len(Presentation("d.pptx").slides))''',
    "odfpy": '''
from odf.opendocument import OpenDocumentText
from odf.text import P
d = OpenDocumentText(); d.text.addElement(P(text="hello odf")); d.save("o.odt")
print("OK", __import__("os").path.getsize("o.odt"))''',
    "fpdf2": '''
from fpdf import FPDF
pdf = FPDF(); pdf.add_page(); pdf.set_font("helvetica", size=14); pdf.cell(text="Reorder list"); pdf.ln(12)
with pdf.table() as t:
    for row in [["Item", "Qty"], ["nuts", "40"]]:
        r = t.row()
        for c in row: r.cell(c)
pdf.output("f.pdf")
print("OK", open("f.pdf", "rb").read()[:8])''',
    "reportlab": '''
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
doc = SimpleDocTemplate("rl.pdf", pagesize=A4)
doc.build([Paragraph("Reorder list", getSampleStyleSheet()["Title"]), Table([["Item", "Qty"], ["nuts", "40"]])])
print("OK", open("rl.pdf", "rb").read()[:8])''',
    "pypdf": '''
import pymupdf
from pypdf import PdfReader, PdfWriter
d = pymupdf.open(); pg = d.new_page(); pg.insert_text((72, 72), "hello pypdf"); d.save("a.pdf"); d.save("b.pdf")
w = PdfWriter(); [w.append(f) for f in ("a.pdf", "b.pdf")]; w.write("ab.pdf")
r = PdfReader("ab.pdf"); print("OK", len(r.pages), r.pages[0].extract_text().strip())''',
    "pdfminer.six": '''
import pymupdf
from pdfminer.high_level import extract_text
d = pymupdf.open(); pg = d.new_page(); pg.insert_text((72, 72), "hello pdfminer"); d.save("m.pdf")
print("OK", extract_text("m.pdf").strip())''',
    "Markdown": '''
import markdown
print("OK", markdown.markdown("# T\\n\\n| a | b |\\n|---|---|\\n| 1 | 2 |", extensions=["tables"])[:60])''',
    "markdown-it-py": '''
from markdown_it import MarkdownIt
print("OK", MarkdownIt().enable("table").render("# T\\n\\n**b**")[:40])''',
    "mistune": '''
import mistune
print("OK", mistune.html("# T\\n\\n*x*")[:40])''',
    "html2text": '''
import html2text
print("OK", html2text.html2text("<h1>T</h1><p>a <b>b</b></p>").strip())''',
    "markdownify": '''
from markdownify import markdownify
print("OK", markdownify("<h1>T</h1><p>a <b>b</b></p>").strip())''',
    "tabulate": '''
from tabulate import tabulate
print("OK\\n" + tabulate([["nuts", 40], ["bolts", 120]], headers=["item", "qty"], tablefmt="github"))''',
    "seaborn": '''
import seaborn as sns, pandas as pd, matplotlib
matplotlib.use("Agg")
ax = sns.barplot(data=pd.DataFrame({"x": ["a", "b"], "y": [1, 3]}), x="x", y="y"); ax.figure.savefig("s.png")
print("OK", __import__("os").path.getsize("s.png"))''',
    "plotly": '''
import plotly.express as px
px.bar(x=["a", "b"], y=[1, 3]).write_html("p.html", include_plotlyjs=True)
print("OK", __import__("os").path.getsize("p.html"))''',
    "xmltodict": '''
import xmltodict
d = xmltodict.parse("<a><b x='1'>t</b></a>"); print("OK", d["a"]["b"]["@x"], xmltodict.unparse(d)[:40])''',
    "icalendar": '''
from icalendar import Calendar, Event
from datetime import datetime
c = Calendar(); e = Event(); e.add("summary", "Meet"); e.add("dtstart", datetime(2026, 10, 6, 9)); c.add_component(e)
print("OK", Calendar.from_ical(c.to_ical()).walk("VEVENT")[0]["summary"])''',
    "feedparser": '''
import feedparser
f = feedparser.parse("<rss version='2.0'><channel><title>T</title><item><title>i1</title></item></channel></rss>")
print("OK", f.feed.title, [e.title for e in f.entries])''',
    "Unidecode": '''
from unidecode import unidecode
print("OK", unidecode("Grüße, Łódź"))''',
    "python-slugify": '''
from slugify import slugify
print("OK", slugify("Grüße aus Köln!"))''',
    "chardet": '''
import chardet
print("OK", chardet.detect("Grüße aus Köln, schöne Straße".encode("latin-1"))["encoding"])''',
    "humanize": '''
import humanize
print("OK", humanize.naturalsize(123456789), humanize.intcomma(1234567))''',
    "num2words": '''
from num2words import num2words
print("OK", num2words(42), num2words(42, lang="de"))''',
    "qrcode": '''
import qrcode
qrcode.make("https://example.org").save("q.png"); print("OK", __import__("os").path.getsize("q.png"))''',
    "segno": '''
import segno
segno.make("hello").save("s.svg"); segno.make("hello").save("s.png"); print("OK", __import__("os").path.getsize("s.png"))''',
    "python-barcode": '''
import barcode
from barcode.writer import ImageWriter
barcode.get("ean13", "590123412345", writer=ImageWriter()).save("bc"); print("OK", sorted(__import__("os").listdir("."))[:3])''',
    "Faker": '''
from faker import Faker
f = Faker("de_DE"); Faker.seed(1); print("OK", f.name(), f.city())''',
}


# What PyPI's metadata doesn't say, found by this smoke test. Bundling has to carry it.
# Distributions needed that the metadata doesn't list (an sdist without metadata):
UNDECLARED_DISTS = {"odfpy": ["defusedxml"]}
# Pyodide packages a library imports at run time but lists only as an optional extra, or
# not at all (by import name, which is what makes the harness load them):
UNDECLARED_PYODIDE = {"fpdf2": ["ssl"], "qrcode": ["PIL"], "python-barcode": ["PIL"], "plotly": ["numpy", "pandas"]}
# Only for the smoke test itself (python-docx makes the input file):
TEST_FIXTURES = {"docx2txt": ["python-docx"], "mammoth": ["python-docx"]}


def norm(name):
    return re.sub(r"[-_.]+", "-", name).lower()


def fetch(url, dest, sha=None):
    if not dest.exists():
        with urllib.request.urlopen(url, timeout=60) as r:
            dest.write_bytes(r.read())
    if sha and hashlib.sha256(dest.read_bytes()).hexdigest() != sha:
        raise RuntimeError(f"sha256 mismatch for {dest.name}")
    return dest


def dist_files(name, cache):
    """The file to bundle for one distribution: its pure wheel, else its sdist."""
    key = norm(name)
    if key not in cache:
        with urllib.request.urlopen(f"https://pypi.org/pypi/{name}/json", timeout=30) as r:
            d = json.load(r)
        pure = [u for u in d["urls"] if u["filename"].endswith("-none-any.whl")]
        sdist = [u for u in d["urls"] if u["packagetype"] == "sdist"]
        u = (pure or sdist)[0]
        cache[key] = fetch(u["url"], WHEELS / u["filename"], u["digests"]["sha256"])
    return cache[key]


def unpack_code(files):
    """Python that unpacks wheels (and sdists) from /workspace into site-packages."""
    return f'''
import zipfile, tarfile, site, os, shutil, importlib
sp = site.getsitepackages()[0]
for f in {files!r}:
    if f.endswith(".whl"):
        zipfile.ZipFile(f).extractall(sp)
    else:   # sdist: copy the top-level packages and modules out of the project folder
        with tarfile.open(f) as t:
            root = t.getnames()[0].split("/")[0]
            tmp = "/tmp/sd-" + root; t.extractall(tmp)
        base = os.path.join(tmp, root)
        for n in os.listdir(base):
            p = os.path.join(base, n)
            if os.path.isdir(p) and os.path.exists(os.path.join(p, "__init__.py")) and n not in ("tests", "test", "docs", "examples"):
                shutil.copytree(p, os.path.join(sp, n), dirs_exist_ok=True)
            elif n.endswith(".py") and n not in ("setup.py", "conftest.py"):
                shutil.copy(p, sp)
importlib.invalidate_caches()
print("unpacked", len({files!r}))
'''


def main():
    study = json.loads(STUDY.read_text())
    lock = json.loads(LOCK.read_text())["packages"]
    imports_of = {norm(v["name"]): v["imports"] for v in lock.values()}
    WHEELS.mkdir(parents=True, exist_ok=True)
    cache, plan = {}, []
    for r in study["rows"]:
        if r.get("error") or r.get("in_pyodide"):
            continue
        extra = [e["name"] for e in r["extra_deps"]] + UNDECLARED_DISTS.get(r["name"], []) + TEST_FIXTURES.get(r["name"], [])
        files = [dist_files(r["name"], cache)] + [dist_files(n, cache) for n in extra]
        # Pyodide packages it needs, by an import name of each, so the harness loads them.
        pre = [imports_of[d][0] for d in r["covered_deps"] if imports_of.get(d)] + UNDECLARED_PYODIDE.get(r["name"], []) + (["lxml", "typing_extensions"] if r["name"] in TEST_FIXTURES else [])
        plan.append((r["name"], files, pre))
    only = [x for x in __import__("os").environ.get("ONLY", "").split(",") if x]
    if only:
        plan = [p for p in plan if p[0] in only]
    all_files = sorted({f for _, fs, _ in plan for f in fs})
    print(f"{len(plan)} libraries, {len(all_files)} files, {sum(f.stat().st_size for f in all_files) / 1e6:.1f} MB")

    results = []
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        for name, files, pre in plan:
            # A fresh interpreter per library: one library's unpacking can't help another.
            pg = b.new_page()
            pg.goto(APP.as_uri())
            while pg.evaluate("() => PY.state") != "idle":
                time.sleep(0.3)
            pg.set_input_files("#wsFileInput", files=[str(f) for f in files])
            while pg.evaluate("() => WS.files.size") < len(files):
                time.sleep(0.2)
            run = lambda code: pg.evaluate("async (c) => { const r = await runInWorker(c, { timeoutMs: 240000 }); return [r.status, r.output]; }", code)  # noqa: E731
            u = run(unpack_code([f.name for f in files]))
            key = next((k for k in SMOKE if norm(k) == norm(name)), None)
            code = ("".join(f"import {i}\n" for i in pre)) + (SMOKE[key] if key else "print('OK (no smoke test)')")
            t0 = time.time()
            st, out = run(code) if u[0] == "ok" else u
            ok = st == "ok" and "OK" in out
            line = (out.strip().splitlines() or [""])[-1][:110]
            print(f"{'✅' if ok else '❌'} {name:16} {time.time() - t0:5.1f} s  {line}")
            results.append({"name": name, "ok": ok, "secs": round(time.time() - t0, 1), "output": out[-1500:], "files": [f.name for f in files], "pyodide_deps": pre})
            pg.close()
        b.close()
    if only and OUT.exists():   # merge into the full run's results
        old = {r["name"]: r for r in json.loads(OUT.read_text())["results"]}
        old.update({r["name"]: r for r in results})
        results = list(old.values())
    OUT.write_text(json.dumps({"date": time.strftime("%Y-%m-%d"), "results": results}, indent=1))
    print(f"\n{sum(r['ok'] for r in results)}/{len(results)} work → {OUT}")


if __name__ == "__main__":
    main()
